from __future__ import annotations

import json
import ipaddress
import re
import socket
import threading
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import tkinter as tk
from tkinter import messagebox, ttk

from openpyxl import load_workbook
from trafilatura import extract


APP_DIR = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = APP_DIR.parent / "liste_entreprises_arsp.xlsx"
DEFAULT_SEARXNG_URL = "http://127.0.0.1:8080"
RESULT_LIMIT = 30
ANALYSIS_LIMIT = 5
MAX_PAGE_BYTES = 2_000_000


@dataclass(frozen=True)
class PageAssessment:
    verdict: str
    score: int
    source_kind: str
    reasons: tuple[str, ...]
    excerpt: str
    final_url: str
    error: str = ""


@dataclass(frozen=True)
class Company:
    legal_name: str
    trade_name: str
    sector: str
    province: str
    rccm: str
    dossier_number: str

    @property
    def label(self) -> str:
        details = [value for value in (self.province, self.sector) if value]
        suffix = f" — {' | '.join(details)}" if details else ""
        return f"{self.legal_name}{suffix} [{self.dossier_number}]"


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def load_companies(workbook_path: Path = DEFAULT_WORKBOOK) -> list[Company]:
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Fichier ARSP introuvable : {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["Données"]
        rows = sheet.iter_rows(values_only=True)
        headers = [_cell_text(value) for value in next(rows)]
        positions = {header: index for index, header in enumerate(headers)}
        required = {"Dénomination", "Numéro Dossier"}
        missing = sorted(required - positions.keys())
        if missing:
            raise ValueError(f"Colonne(s) ARSP manquante(s) : {', '.join(missing)}")

        def value(row: tuple[Any, ...], header: str) -> str:
            index = positions.get(header)
            return _cell_text(row[index]) if index is not None and index < len(row) else ""

        companies = []
        for row in rows:
            legal_name = value(row, "Dénomination")
            dossier_number = value(row, "Numéro Dossier")
            if not legal_name or not dossier_number:
                continue
            companies.append(
                Company(
                    legal_name=legal_name,
                    trade_name=value(row, "Nom Commercial"),
                    sector=value(row, "Activité"),
                    province=value(row, "Province"),
                    rccm=value(row, "RCCM"),
                    dossier_number=dossier_number,
                )
            )
        return sorted(companies, key=lambda company: company.legal_name.casefold())
    finally:
        workbook.close()


def build_query(company_name: str, sector: str = "", keywords: str = "", rccm: str = "") -> str:
    company_name = company_name.strip()
    if not company_name:
        raise ValueError("Choisis ou saisis le nom d’une entreprise.")

    parts = [f'"{company_name}"']
    if rccm.strip():
        parts.append(f'"{rccm.strip()}"')
    if sector.strip():
        parts.append(sector.strip())
    if keywords.strip():
        parts.append(keywords.strip())
    return " ".join(parts)


def _normalise_for_match(value: str) -> str:
    value = unicodedata.normalize("NFKD", value.casefold())
    value = "".join(character for character in value if not unicodedata.combining(character))
    return " ".join(re.findall(r"[a-z0-9]+", value))


def _meaningful_tokens(value: str) -> set[str]:
    ignored = {"avec", "dans", "pour", "sarl", "sarlu", "sprl", "sa", "sas", "societe", "entreprise"}
    return {token for token in _normalise_for_match(value).split() if len(token) >= 4 and token not in ignored}


def source_kind(url: str, company: Company) -> str:
    host = (urllib.parse.urlsplit(url).hostname or "").casefold().removeprefix("www.")
    if host.endswith((".gouv.cd", ".gov.cd")) or host in {"arsp.cd", "www.arsp.cd", "guichetunique.cd"}:
        return "Institutionnelle"
    if host.endswith(("facebook.com", "linkedin.com", "instagram.com", "x.com", "twitter.com", "youtube.com")):
        return "Réseau social"
    if host.endswith(("scribd.com", "goafricaonline.com", "africabizinfo.com", "pagesjaunes.cd")):
        return "Annuaire / plateforme"
    host_words = set(re.findall(r"[a-z0-9]+", host))
    if host_words & _meaningful_tokens(company.legal_name):
        return "Site d’entreprise possible"
    return "Web"


def assess_page(company: Company, title: str, text: str, url: str, keywords: str = "") -> PageAssessment:
    searchable = _normalise_for_match(f"{title}\n{text}")
    searchable_words = set(searchable.split())
    reasons: list[str] = []
    score = 0

    legal_name = _normalise_for_match(company.legal_name)
    trade_name = _normalise_for_match(company.trade_name)
    rccm = _normalise_for_match(company.rccm)
    legal_match = bool(legal_name and legal_name in searchable)
    trade_match = bool(trade_name and len(trade_name) >= 5 and trade_name in searchable)
    rccm_match = bool(rccm and rccm in searchable)

    if legal_match:
        score += 5
        reasons.append("nom légal trouvé dans la page")
    elif trade_match:
        score += 3
        reasons.append("nom commercial trouvé dans la page")

    if rccm_match:
        score += 5
        reasons.append("RCCM trouvé dans la page")

    province = _normalise_for_match(company.province)
    if province and province in searchable:
        score += 1
        reasons.append("province ou ville concordante")

    sector_tokens = _meaningful_tokens(company.sector)
    matched_sector = sector_tokens & searchable_words
    if sector_tokens and len(matched_sector) >= max(1, len(sector_tokens) // 2):
        score += 2
        reasons.append("activité concordante")
    elif matched_sector:
        score += 1
        reasons.append("activité partiellement concordante")

    keyword_tokens = _meaningful_tokens(keywords)
    if keyword_tokens & searchable_words:
        score += 1
        reasons.append("mot-clé supplémentaire trouvé")

    if rccm_match or (legal_match and score >= 6):
        verdict = "Pertinente"
    elif legal_match or (trade_match and score >= 4):
        verdict = "À vérifier"
    else:
        verdict = "Hors sujet probable"
        reasons.append("pas assez d’éléments pour relier cette page à l’entreprise")

    excerpt = relevant_excerpt(text, company, keywords)
    return PageAssessment(
        verdict=verdict,
        score=score,
        source_kind=source_kind(url, company),
        reasons=tuple(reasons),
        excerpt=excerpt,
        final_url=url,
    )


def relevant_excerpt(text: str, company: Company, keywords: str = "", limit: int = 900) -> str:
    sentences = [sentence.strip() for sentence in re.split(r"(?<=[.!?])\s+|\n+", text) if sentence.strip()]
    wanted = _meaningful_tokens(" ".join((company.legal_name, company.trade_name, company.sector, company.province, keywords)))
    ranked = sorted(
        enumerate(sentences),
        key=lambda item: (len(_meaningful_tokens(item[1]) & wanted), -item[0]),
        reverse=True,
    )
    chosen = [sentence for _index, sentence in ranked[:3] if _meaningful_tokens(sentence) & wanted]
    excerpt = "\n\n".join(chosen) if chosen else "Aucun passage clairement pertinent n’a été extrait."
    return excerpt[:limit].rstrip()


def validate_public_url(url: str) -> None:
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("URL non HTTP refusée")
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    try:
        addresses = socket.getaddrinfo(parsed.hostname, port, type=socket.SOCK_STREAM)
    except socket.gaierror as error:
        raise ValueError("nom de domaine introuvable") from error
    if not addresses:
        raise ValueError("nom de domaine sans adresse")
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if not ip.is_global:
            raise ValueError("adresse locale ou privée refusée")


class _PublicRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, file_pointer, code, message, headers, new_url):
        validate_public_url(new_url)
        return super().redirect_request(request, file_pointer, code, message, headers, new_url)


class PageAnalyzer:
    def __init__(self, timeout: int = 12) -> None:
        self.timeout = timeout
        self.opener = urllib.request.build_opener(_PublicRedirectHandler())

    def analyze(self, result: dict[str, Any], company: Company, keywords: str = "") -> PageAssessment:
        url = str(result.get("url") or "")
        title = str(result.get("title") or "Sans titre")
        try:
            validate_public_url(url)
            request = urllib.request.Request(
                url,
                headers={
                    "User-Agent": "Mozilla/5.0 (compatible; Onbora-Essai/0.2; local prototype)",
                    "Accept": "text/html,application/xhtml+xml",
                },
            )
            with self.opener.open(request, timeout=self.timeout) as response:
                final_url = response.geturl()
                validate_public_url(final_url)
                content_type = response.headers.get_content_type()
                if content_type not in {"text/html", "application/xhtml+xml"}:
                    raise ValueError(f"format non analysé : {content_type}")
                raw = response.read(MAX_PAGE_BYTES + 1)
                if len(raw) > MAX_PAGE_BYTES:
                    raise ValueError("page trop volumineuse")
                charset = response.headers.get_content_charset() or "utf-8"
                html = raw.decode(charset, errors="replace")
            text = extract(
                html,
                url=final_url,
                include_comments=False,
                include_tables=False,
                favor_precision=True,
            )
            if not text or len(text.strip()) < 80:
                raise ValueError("aucun texte principal exploitable")
            return assess_page(company, title, text, final_url, keywords)
        except (ValueError, urllib.error.URLError, urllib.error.HTTPError, TimeoutError, OSError) as error:
            return PageAssessment(
                verdict="Inaccessible",
                score=0,
                source_kind=source_kind(url, company),
                reasons=("la page n’a pas pu être analysée",),
                excerpt="",
                final_url=url,
                error=str(error),
            )


class SearxngClient:
    def __init__(self, base_url: str = DEFAULT_SEARXNG_URL, timeout: int = 25) -> None:
        self.base_url = base_url.rstrip("/")
        self.timeout = timeout

    def search(self, query: str) -> list[dict[str, Any]]:
        params = urllib.parse.urlencode(
            {
                "q": query,
                "format": "json",
                "language": "fr",
                "safesearch": "1",
            }
        )
        request = urllib.request.Request(
            f"{self.base_url}/search?{params}",
            headers={"User-Agent": "Onbora-Essai/0.1"},
        )
        try:
            with urllib.request.urlopen(request, timeout=self.timeout) as response:
                payload = json.load(response)
        except urllib.error.HTTPError as error:
            if error.code == 403:
                raise RuntimeError("SearXNG refuse le format JSON. Vérifie son fichier settings.yml.") from error
            raise RuntimeError(f"SearXNG a répondu avec l’erreur HTTP {error.code}.") from error
        except urllib.error.URLError as error:
            raise RuntimeError(
                "SearXNG n’est pas joignable. Lance d’abord le fichier lancer.ps1."
            ) from error
        except (json.JSONDecodeError, UnicodeDecodeError) as error:
            raise RuntimeError("SearXNG a renvoyé une réponse illisible.") from error

        results = payload.get("results", [])
        if not isinstance(results, list):
            raise RuntimeError("La réponse SearXNG ne contient pas de liste de résultats.")
        return [result for result in results[:RESULT_LIMIT] if isinstance(result, dict)]


class SearchApp:
    def __init__(self, root: tk.Tk, companies: list[Company]) -> None:
        self.root = root
        self.companies = companies
        self.company_by_label = {company.label: company for company in companies}
        self.results: list[dict[str, Any]] = []
        self.client = SearxngClient()
        self.page_analyzer = PageAnalyzer()

        root.title("Onbora — Essai de recherche")
        root.geometry("1100x720")
        root.minsize(820, 560)

        self.company_var = tk.StringVar()
        self.sector_var = tk.StringVar()
        self.keywords_var = tk.StringVar()
        self.company_info_var = tk.StringVar(value=f"{len(companies):,} entreprises ARSP chargées".replace(",", " "))
        self.status_var = tk.StringVar(value="Prêt à rechercher avec SearXNG local.")

        self._build_ui()

    def _build_ui(self) -> None:
        shell = ttk.Frame(self.root, padding=18)
        shell.pack(fill="both", expand=True)
        shell.columnconfigure(0, weight=1)
        shell.rowconfigure(2, weight=1)

        ttk.Label(shell, text="Recherche simple d’entreprise", font=("Segoe UI", 18, "bold")).grid(
            row=0, column=0, sticky="w"
        )
        ttk.Label(
            shell,
            text="Choisis une entreprise ARSP, complète la recherche si nécessaire, puis regarde les résultats.",
        ).grid(row=1, column=0, sticky="w", pady=(4, 16))

        content = ttk.Panedwindow(shell, orient="vertical")
        content.grid(row=2, column=0, sticky="nsew")

        form = ttk.LabelFrame(content, text="Recherche", padding=14)
        form.columnconfigure(1, weight=1)
        content.add(form, weight=0)

        ttk.Label(form, text="Entreprise").grid(row=0, column=0, sticky="w", padx=(0, 12), pady=5)
        self.company_combo = ttk.Combobox(
            form,
            textvariable=self.company_var,
            values=[company.label for company in self.companies],
            state="normal",
        )
        self.company_combo.grid(row=0, column=1, sticky="ew", pady=5)
        self.company_combo.bind("<<ComboboxSelected>>", self._company_selected)
        self.company_combo.focus_set()

        ttk.Label(form, text="Secteur").grid(row=1, column=0, sticky="w", padx=(0, 12), pady=5)
        ttk.Entry(form, textvariable=self.sector_var).grid(row=1, column=1, sticky="ew", pady=5)

        ttk.Label(form, text="Mots-clés / lieu").grid(row=2, column=0, sticky="w", padx=(0, 12), pady=5)
        keywords_entry = ttk.Entry(form, textvariable=self.keywords_var)
        keywords_entry.grid(row=2, column=1, sticky="ew", pady=5)
        keywords_entry.bind("<Return>", lambda _event: self.start_search())

        actions = ttk.Frame(form)
        actions.grid(row=3, column=1, sticky="ew", pady=(10, 0))
        actions.columnconfigure(0, weight=1)
        ttk.Label(actions, textvariable=self.company_info_var).grid(row=0, column=0, sticky="w")
        self.search_button = ttk.Button(actions, text="Rechercher", command=self.start_search)
        self.search_button.grid(row=0, column=1, padx=(8, 0))

        results_frame = ttk.LabelFrame(content, text="Résultats et analyse locale", padding=10)
        results_frame.columnconfigure(0, weight=1)
        results_frame.rowconfigure(0, weight=3)
        results_frame.rowconfigure(1, weight=1)
        content.add(results_frame, weight=1)

        self.result_tree = ttk.Treeview(
            results_frame,
            columns=("verdict", "kind", "title"),
            show="headings",
            selectmode="browse",
        )
        self.result_tree.heading("verdict", text="Pertinence")
        self.result_tree.heading("kind", text="Type de source")
        self.result_tree.heading("title", text="Titre")
        self.result_tree.column("verdict", width=140, anchor="w", stretch=False)
        self.result_tree.column("kind", width=180, anchor="w", stretch=False)
        self.result_tree.column("title", width=650, anchor="w")
        tree_scroll = ttk.Scrollbar(results_frame, orient="vertical", command=self.result_tree.yview)
        self.result_tree.configure(yscrollcommand=tree_scroll.set)
        self.result_tree.grid(row=0, column=0, sticky="nsew")
        tree_scroll.grid(row=0, column=1, sticky="ns")
        self.result_tree.bind("<<TreeviewSelect>>", self._show_selected_result)
        self.result_tree.bind("<Double-1>", lambda _event: self.open_selected_result())

        detail_frame = ttk.Frame(results_frame)
        detail_frame.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=(10, 0))
        detail_frame.columnconfigure(0, weight=1)
        detail_frame.rowconfigure(0, weight=1)
        self.detail_text = tk.Text(detail_frame, height=7, wrap="word", state="disabled")
        self.detail_text.grid(row=0, column=0, sticky="nsew")
        detail_actions = ttk.Frame(detail_frame)
        detail_actions.grid(row=1, column=0, sticky="ew", pady=(8, 0))
        detail_actions.columnconfigure(0, weight=1)
        self.analyze_button = ttk.Button(
            detail_actions,
            text=f"Analyser les {ANALYSIS_LIMIT} premières pages",
            command=self.start_analysis,
            state="disabled",
        )
        self.analyze_button.grid(row=0, column=0, sticky="w")
        ttk.Button(detail_actions, text="Ouvrir le lien", command=self.open_selected_result).grid(
            row=0, column=1, sticky="e"
        )

        ttk.Label(shell, textvariable=self.status_var, anchor="w").grid(row=3, column=0, sticky="ew", pady=(10, 0))

    def _company_selected(self, _event: tk.Event | None = None) -> None:
        company = self.company_by_label.get(self.company_var.get())
        if not company:
            return
        self.sector_var.set(company.sector)
        self.keywords_var.set(company.province)
        info = [f"Dossier : {company.dossier_number}"]
        if company.rccm:
            info.append(f"RCCM : {company.rccm}")
        self.company_info_var.set(" · ".join(info))

    def start_search(self) -> None:
        company = self.company_by_label.get(self.company_var.get())
        company_name = company.legal_name if company else self.company_var.get()
        try:
            query = build_query(company_name, self.sector_var.get(), self.keywords_var.get())
        except ValueError as error:
            messagebox.showwarning("Recherche incomplète", str(error), parent=self.root)
            return

        self.search_button.configure(state="disabled")
        self.status_var.set(f"Recherche en cours : {query}")
        self._clear_results()
        threading.Thread(target=self._search_in_background, args=(query,), daemon=True).start()

    def _search_in_background(self, query: str) -> None:
        try:
            results = self.client.search(query)
        except Exception as error:
            self.root.after(0, self._search_failed, str(error))
            return
        self.root.after(0, self._display_results, results)

    def _display_results(self, results: list[dict[str, Any]]) -> None:
        self.results = results
        for index, result in enumerate(results):
            title = str(result.get("title") or "Sans titre")
            self.result_tree.insert("", "end", iid=str(index), values=("Non analysée", "—", title))
        self.search_button.configure(state="normal")
        self.analyze_button.configure(state="normal" if results else "disabled")
        if results:
            self.status_var.set(f"{len(results)} résultat(s) affiché(s). Double-clique pour ouvrir un lien.")
            self.result_tree.selection_set("0")
            self.result_tree.focus("0")
            self._show_selected_result()
        else:
            self.status_var.set("Aucun résultat. Essaie avec moins de mots-clés.")

    def _search_failed(self, message: str) -> None:
        self.search_button.configure(state="normal")
        self.analyze_button.configure(state="disabled")
        self.status_var.set("La recherche a échoué.")
        messagebox.showerror("Recherche impossible", message, parent=self.root)

    def _clear_results(self) -> None:
        self.results = []
        self.analyze_button.configure(state="disabled")
        for item in self.result_tree.get_children():
            self.result_tree.delete(item)
        self._set_detail("")

    def _selected_result(self) -> dict[str, Any] | None:
        selection = self.result_tree.selection()
        if not selection:
            return None
        index = int(selection[0])
        return self.results[index] if index < len(self.results) else None

    def _show_selected_result(self, _event: tk.Event | None = None) -> None:
        result = self._selected_result()
        if not result:
            self._set_detail("")
            return
        url = str(result.get("url") or "")
        content = str(result.get("content") or "Aucun résumé disponible.")
        assessment = result.get("assessment")
        if isinstance(assessment, PageAssessment):
            reasons = " ; ".join(assessment.reasons)
            if assessment.error:
                body = f"{assessment.verdict} — {assessment.source_kind}\n{reasons}\nErreur : {assessment.error}"
            else:
                body = (
                    f"{assessment.verdict} — {assessment.source_kind} — score {assessment.score}\n"
                    f"{reasons}\n\nPassages extraits :\n{assessment.excerpt}"
                )
            self._set_detail(f"{assessment.final_url}\n\n{body}")
        else:
            self._set_detail(f"{url}\n\nRésumé du moteur de recherche (page non encore analysée) :\n{content}")

    def start_analysis(self) -> None:
        if not self.results:
            return
        selected_company = self.company_by_label.get(self.company_var.get())
        if selected_company:
            company = selected_company
        else:
            company_name = self.company_var.get().strip()
            if not company_name:
                return
            company = Company(
                legal_name=company_name,
                trade_name="",
                sector=self.sector_var.get().strip(),
                province="",
                rccm="",
                dossier_number="saisie libre",
            )
        self.search_button.configure(state="disabled")
        self.analyze_button.configure(state="disabled")
        self.status_var.set(f"Analyse locale des {min(ANALYSIS_LIMIT, len(self.results))} premières pages...")
        threading.Thread(
            target=self._analyze_in_background,
            args=(company, self.keywords_var.get()),
            daemon=True,
        ).start()

    def _analyze_in_background(self, company: Company, keywords: str) -> None:
        assessments: list[tuple[int, PageAssessment]] = []
        for index, result in enumerate(self.results[:ANALYSIS_LIMIT]):
            assessment = self.page_analyzer.analyze(result, company, keywords)
            assessments.append((index, assessment))
            self.root.after(
                0,
                self.status_var.set,
                f"Page {index + 1}/{min(ANALYSIS_LIMIT, len(self.results))} analysée : {assessment.verdict}",
            )
        self.root.after(0, self._display_assessments, assessments)

    def _display_assessments(self, assessments: list[tuple[int, PageAssessment]]) -> None:
        for index, assessment in assessments:
            self.results[index]["assessment"] = assessment
            if self.result_tree.exists(str(index)):
                self.result_tree.set(str(index), "verdict", assessment.verdict)
                self.result_tree.set(str(index), "kind", assessment.source_kind)
        self.search_button.configure(state="normal")
        self.analyze_button.configure(state="normal")
        relevant = sum(assessment.verdict == "Pertinente" for _index, assessment in assessments)
        uncertain = sum(assessment.verdict == "À vérifier" for _index, assessment in assessments)
        inaccessible = sum(assessment.verdict == "Inaccessible" for _index, assessment in assessments)
        self.status_var.set(
            f"Analyse terminée : {relevant} pertinente(s), {uncertain} à vérifier, "
            f"{inaccessible} inaccessible(s)."
        )
        self._show_selected_result()

    def _set_detail(self, text: str) -> None:
        self.detail_text.configure(state="normal")
        self.detail_text.delete("1.0", "end")
        self.detail_text.insert("1.0", text)
        self.detail_text.configure(state="disabled")

    def open_selected_result(self) -> None:
        result = self._selected_result()
        if not result:
            return
        url = str(result.get("url") or "")
        if url.startswith(("http://", "https://")):
            webbrowser.open(url)


def main() -> None:
    try:
        companies = load_companies()
    except Exception as error:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("Démarrage impossible", str(error), parent=root)
        root.destroy()
        raise SystemExit(1) from error

    root = tk.Tk()
    SearchApp(root, companies)
    root.mainloop()


if __name__ == "__main__":
    main()
