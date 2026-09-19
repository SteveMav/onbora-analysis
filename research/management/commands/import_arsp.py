"""Import an ARSP workbook without mutating values used as provenance."""

from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from research.models import ArspImport, ArspRecord
from research.text_utils import clean_excel_text

REQUIRED_HEADERS = {"Numéro Dossier", "Dénomination"}
EXPECTED_HEADERS = {
    "Dénomination", "Nom Commercial", "RCCM", "Numéro Dossier", "Activité",
    "Date Expiration", "Province", "Num Impôt", "Adresse", "Téléphone", "Email",
}
SPACE_RE = re.compile(r"\s+")


def _as_json_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _normalise(value: str) -> str:
    return clean_excel_text(value).casefold()


class Command(BaseCommand):
    help = "Import and validate an ARSP XLSX file as an immutable import version."

    def add_arguments(self, parser) -> None:
        parser.add_argument("workbook", type=Path)
        parser.add_argument("--sheet", default="Données")

    def handle(self, *args, **options) -> None:
        source: Path = options["workbook"]
        if not source.is_file():
            raise CommandError("Workbook not found")
        file_hash = hashlib.sha256(source.read_bytes()).hexdigest()
        if ArspImport.objects.filter(file_hash=file_hash).exists():
            self.stdout.write(self.style.WARNING("This exact workbook is already imported; no new version created."))
            return

        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook[options["sheet"]]
        except KeyError as error:
            raise CommandError(f"Worksheet {options['sheet']!r} was not found") from error

        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [_as_json_value(value) for value in next(rows)]
        except StopIteration as error:
            raise CommandError("Workbook is empty") from error
        if len(headers) != len(set(headers)):
            raise CommandError("Workbook has duplicate headers")
        missing = REQUIRED_HEADERS - set(headers)
        if missing:
            raise CommandError(f"Workbook lacks required header(s): {', '.join(sorted(missing))}")

        dossier_index = headers.index("Numéro Dossier")
        legal_name_index = headers.index("Dénomination")
        records: list[tuple[dict[str, str], dict[str, str], list[str]]] = []
        dossiers: set[str] = set()
        errors: list[str] = []
        whitespace_counts = {header: 0 for header in headers}
        missing_counts = {header: 0 for header in headers}
        for row_number, row in enumerate(rows, start=2):
            raw = {header: _as_json_value(row[index] if index < len(row) else None) for index, header in enumerate(headers)}
            normalized = {header: _normalise(value) for header, value in raw.items()}
            warnings: list[str] = []
            for header, value in raw.items():
                if not value:
                    missing_counts[header] += 1
                elif value != SPACE_RE.sub(" ", value).strip():
                    whitespace_counts[header] += 1
                    warnings.append(f"normalization_available:{header}")
            dossier = normalized["Numéro Dossier"]
            legal_name = clean_excel_text(raw["Dénomination"])
            if not dossier or not legal_name:
                errors.append(f"row {row_number}: Numéro Dossier and Dénomination are required")
            elif dossier in dossiers:
                errors.append(f"row {row_number}: duplicate Numéro Dossier")
            dossiers.add(dossier)
            records.append((raw, normalized, warnings))
        workbook.close()

        if errors:
            report = {"row_count": len(records), "errors": errors[:50], "error_count": len(errors)}
            raise CommandError(json.dumps(report, ensure_ascii=False))

        report = {
            "row_count": len(records),
            "headers": headers,
            "unexpected_headers": sorted(set(headers) - EXPECTED_HEADERS),
            "missing_counts": missing_counts,
            "whitespace_normalization_counts": whitespace_counts,
            "error_count": 0,
        }
        with transaction.atomic():
            arsp_import = ArspImport.objects.create(
                file_name=source.name,
                file_hash=file_hash,
                validation_report=report,
            )
            ArspRecord.objects.bulk_create([
                ArspRecord(
                    arsp_import=arsp_import,
                    dossier_number=normalized["Numéro Dossier"],
                    legal_name=clean_excel_text(raw["Dénomination"]),
                    raw_data=raw,
                    normalized_data=normalized,
                    warnings=warnings,
                )
                for raw, normalized, warnings in records
            ], batch_size=500)
        self.stdout.write(self.style.SUCCESS(json.dumps(report, ensure_ascii=False)))
